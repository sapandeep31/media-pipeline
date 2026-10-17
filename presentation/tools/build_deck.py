"""Build out/deck.pptx from a PPTX template + a content file.

A deck project is a folder (default: deck/, or set DECK_DIR) holding template.pptx and content.py;
image paths in content.py resolve against that folder first, then the presentation/ root.

Every build starts from a fresh copy of the template, so the output is reproducible:
  1. structure: duplicate / delete / re-order template slides (content.ORDER)
  2. cover: fill Team Details (content.COVER)
  3. per slide (content.SLIDES): keep or rename the template's section label, add the action title,
     add elements (text, image, rect, bar chart), speaker notes

Usage: .venv\\Scripts\\python.exe tools\\build_deck.py [--out out\\deck.pptx]   (set DECK_DIR=projects\\x to switch project)
"""
import argparse
import copy
import importlib.util
import os
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DECK_DIR = Path(os.environ.get("DECK_DIR", ROOT / "deck")).resolve()
TEMPLATE = DECK_DIR / "template.pptx"
RT_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"

spec = importlib.util.spec_from_file_location("content", DECK_DIR / "content.py")
C = importlib.util.module_from_spec(spec)
spec.loader.exec_module(C)

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


# ---------------------------------------------------------------- structure
def duplicate(prs, src):
    """Clone a template slide (shapes + image rels) at the end of the deck."""
    new = prs.slides.add_slide(src.slide_layout)
    for ph in list(new.placeholders):
        ph._element.getparent().remove(ph._element)
    rid_map = {}
    for rel in src.part.rels.values():
        if rel.reltype == RT_IMAGE:
            rid_map[rel.rId] = new.part.relate_to(rel._target, RT_IMAGE)
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        clone = copy.deepcopy(el)
        for node in clone.iter():
            for attr in (qn("r:embed"), qn("r:link")):
                if node.get(attr) in rid_map:
                    node.set(attr, rid_map[node.get(attr)])
        new.shapes._spTree.append(clone)
    return new


def restructure(prs):
    """content.ORDER lists template slide numbers (1-based); a number may repeat (= duplicate)."""
    originals = list(prs.slides)
    ids = prs.slides._sldIdLst
    orig_els = list(ids)                          # same order as originals
    used, order_els, slides = set(), [], []
    for n in C.ORDER:
        src = originals[n - 1]
        if n in used:
            s = duplicate(prs, src)
            el = ids[-1]
        else:
            s, el = src, orig_els[n - 1]
            used.add(n)
        order_els.append(el)
        slides.append((n, s))
    for el in list(ids):
        ids.remove(el)
    for el in order_els:
        ids.append(el)
    for n, el in enumerate(orig_els, 1):          # drop slides not in ORDER
        if n not in used:
            prs.part.drop_rel(el.get(qn("r:id")))
    return slides


# ---------------------------------------------------------------- text
def _run_style(run, st):
    f = run.font
    name = st.get("font", C.FONT)
    if name == C.FONT_SEMI:
        # portable: regular family + bold flag renders bold everywhere (PowerPoint, Google Slides, LibreOffice),
        # even where the "SemiBold" family name isn't installed
        name, st = C.FONT, {**st, "bold": True}
    f.name = name
    f.size = Pt(st.get("size", C.BODY_PT))
    if "bold" in st:
        f.bold = st["bold"]
    if st.get("italic"):
        f.italic = True
    f.color.rgb = RGBColor.from_string(st.get("color", C.INK))


def _bullet(p, st):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(st.get("indent", 0.22)))))
    pPr.set("indent", str(-int(Inches(st.get("indent", 0.22)))))
    for tag in ("a:buNone", "a:buChar", "a:buFont", "a:buClr", "a:buSzPct"):
        for old in pPr.findall(qn(tag)):
            pPr.remove(old)
    clr = pPr.makeelement(qn("a:buClr"), {})
    clr.append(clr.makeelement(qn("a:srgbClr"), {"val": st.get("bullet_color", C.TEAL_FILL)}))
    pPr.append(clr)
    pPr.append(pPr.makeelement(qn("a:buSzPct"), {"val": "100000"}))
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))  # XML bullet glyph, not literal text


def add_text(slide, e):
    """e: x,y,w,h (in), paras=[str | [(text, style), ...] | {"runs":..., "bullet":True, ...}], style keys."""
    tb = slide.shapes.add_textbox(Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
    tb.name = e.get("name", "text")
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    m = e.get("margin", 0)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(m)
    tf.vertical_anchor = ANCHOR[e.get("anchor", "t")]
    base = {k: v for k, v in e.items() if k in ("font", "size", "bold", "color", "italic")}
    for i, para in enumerate(e["paras"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pst = dict(base)
        if isinstance(para, dict):
            pst.update({k: v for k, v in para.items() if k != "runs"})
            runs = para["runs"]
        else:
            runs = para
        if isinstance(runs, str):
            runs = [(runs, {})]
        p.alignment = ALIGN[pst.get("align", e.get("align", "l"))]
        if "line" in pst or "line" in e:
            p.line_spacing = pst.get("line", e.get("line"))
        p.space_after = Pt(pst.get("after", e.get("after", 0)))
        if pst.get("bullet"):
            _bullet(p, pst)
        for text, rst in runs:
            r = p.add_run()
            r.text = text
            _run_style(r, {**pst, **rst})
            if text.startswith("http"):            # clickable in the PPTX and the exported PDF
                r.hyperlink.address = text
    return tb


# ---------------------------------------------------------------- image / shapes / chart
def add_image(slide, e):
    """Fill the box (x,y,w,h in) like CSS object-fit: cover, around focus (fx, fy in 0..1).
    Optional src_box=(l,t,r,b) in source pixels to pre-crop."""
    path = DECK_DIR / e["src"] if (DECK_DIR / e["src"]).exists() else ROOT / e["src"]
    with Image.open(path) as im:
        W, H = im.size
    l, t, r, b = e.get("src_box", (0, 0, W, H))
    bw, bh = r - l, b - t
    box_ar, img_ar = e["w"] / e["h"], bw / bh
    if img_ar > box_ar:                       # too wide: trim the sides
        cw, ch = bh * box_ar, bh
    else:                                     # too tall: trim top/bottom
        cw, ch = bw, bw / box_ar
    fx, fy = e.get("focus", (0.5, 0.5))
    cl = l + (bw - cw) * fx
    ct = t + (bh - ch) * fy
    # bake the crop and downscale to ~220 dpi at placed size: small file, no reliance on viewer crop support
    max_w = int(e["w"] * 220)
    out = ROOT / "out/img" / f"{path.stem}-{int(cl)}-{int(ct)}-{int(cw)}-{int(ch)}-{max_w}.jpg"
    if not out.exists():
        out.parent.mkdir(parents=True, exist_ok=True)
        with Image.open(path) as im:
            crop = im.convert("RGB").crop((int(cl), int(ct), int(cl + cw), int(ct + ch)))
            if crop.width > max_w:
                crop = crop.resize((max_w, round(max_w * crop.height / crop.width)), Image.LANCZOS)
            crop.save(out, "JPEG", quality=90, optimize=True)
    pic = slide.shapes.add_picture(str(out), Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
    pic.name = e.get("name", path.stem)
    if e.get("border"):
        pic.line.color.rgb = RGBColor.from_string(e["border"])
        pic.line.width = Pt(e.get("border_pt", 0.75))
    return pic


def add_rect(slide, e):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if e.get("radius") else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape, Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
    s.name = e.get("name", "rect")
    style = s._element.find(qn("p:style"))       # theme style adds a soft shadow in LibreOffice/Slides
    if style is not None:
        s._element.remove(style)
    if e.get("radius"):
        s.adjustments[0] = e["radius"]
    if e.get("fill"):
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(e["fill"])
    else:
        s.fill.background()
    if e.get("line"):
        s.line.color.rgb = RGBColor.from_string(e["line"])
        s.line.width = Pt(e.get("line_pt", 1))
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if e.get("text"):
        tf = s.text_frame
        tf.margin_left = tf.margin_right = Inches(e.get("pad", 0.08))
        tf.margin_top = tf.margin_bottom = Inches(0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = ALIGN[e.get("align", "c")]
        r = p.add_run()
        r.text = e["text"]
        _run_style(r, {k: v for k, v in e.items() if k in ("font", "size", "bold", "color")})
    return s


def add_hbar(slide, e):
    """Horizontal bar chart, one series, decluttered, direct value labels, per-bar colours."""
    data = CategoryChartData()
    data.categories = [c for c, _ in e["bars"]]
    data.add_series(e.get("series", "value"), [v for _, v in e["bars"]])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(e["x"]), Inches(e["y"]),
                                Inches(e["w"]), Inches(e["h"]), data)
    gf.name = e.get("name", "chart")
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name = C.FONT
    ch.font.size = Pt(e.get("size", 11))
    ch.font.color.rgb = RGBColor.from_string(C.INK)
    plot = ch.plots[0]
    plot.gap_width = e.get("gap", 45)
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = e.get("fmt", "0.0")
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(e.get("label_size", 12))
    dl.font.bold = True
    dl.font.name = C.FONT_SEMI
    cat = ch.category_axis
    cat.reverse_order = True                   # first bar on top
    cat.has_major_gridlines = False
    cat.format.line.fill.background()
    cat.tick_labels.font.size = Pt(e.get("size", 11))
    val = ch.value_axis
    val.visible = False
    val.has_major_gridlines = False
    val.minimum_scale = 0                      # zero baseline, always
    if e.get("max"):
        val.maximum_scale = e["max"]
    for i, pt in enumerate(plot.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = RGBColor.from_string(e["colors"][i % len(e["colors"])])
    return gf


def add_icon(slide, e):
    """Tabler outline icon (name or path), recoloured, rasterised once to out/icons/<name>-<color>.png."""
    import pymupdf
    name = e["name"]
    svg_path = ROOT / "assets/icons" / f"{name}.svg"
    if not svg_path.exists():
        svg_path = ROOT / "node_modules/@tabler/icons/icons/outline" / f"{name}.svg"
    color = e.get("color", C.TEAL_TEXT)
    png = ROOT / "out/icons" / f"{name}-{color}.png"
    if not png.exists():
        png.parent.mkdir(parents=True, exist_ok=True)
        svg = svg_path.read_text(encoding="utf-8").replace("currentColor", f"#{color}")
        svg = svg.replace('stroke-width="2"', f'stroke-width="{e.get("stroke", 1.75)}"')
        doc = pymupdf.open(stream=svg.encode(), filetype="svg")
        doc[0].get_pixmap(matrix=pymupdf.Matrix(12, 12), alpha=True).save(png)  # 24 px viewBox -> 288 px
    s = e.get("size", 0.32)
    pic = slide.shapes.add_picture(str(png), Inches(e["x"]), Inches(e["y"]), Inches(s), Inches(s))
    pic.name = f"icon-{name}"
    return pic


def add_arrow(slide, e):
    """Straight connector with an arrowhead (x1,y1 -> x2,y2 in inches)."""
    from pptx.enum.shapes import MSO_CONNECTOR
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(e["x1"]), Inches(e["y1"]), Inches(e["x2"]), Inches(e["y2"]))
    c.line.color.rgb = RGBColor.from_string(e.get("color", C.INK_MUTED))
    c.line.width = Pt(e.get("pt", 1.25))
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    c.name = e.get("name", "arrow")
    return c


ADD = {"text": add_text, "image": add_image, "rect": add_rect, "hbar": add_hbar, "icon": add_icon, "arrow": add_arrow}


# ---------------------------------------------------------------- template heading
def heading_shape(slide):
    for sh in slide.shapes:
        if not sh.name.startswith("ms-") and sh.has_text_frame and sh.text_frame.text.strip():
            return sh
    return None


def set_heading(slide, label, move=True):
    """Keep the template's section label (its run formatting), drop prompt lines, shrink the box."""
    sh = heading_shape(slide)
    if sh is None:
        if label:  # blank template slide: add a label styled like the template's headings
            add_text(slide, {"name": "ms-section-label", "x": 0.385, "y": C.LABEL_Y, "w": 9.2, "h": C.LABEL_H,
                             "paras": [label], "font": "Google Sans Flex SemiBold", "size": 16, "color": "000000"})
        return
    paras = sh.text_frame.paragraphs
    first = paras[0]
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    for br in first._p.findall(qn("a:br")):
        first._p.remove(br)
    runs = first.runs
    if label is not None and runs:
        runs[0].text = label
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    if move:  # content slides: pin the label to the standard position; covers keep the template's layout
        sh.top, sh.height = Inches(C.LABEL_Y), Inches(C.LABEL_H)
        tf = sh.text_frame
        tf.margin_top = tf.margin_bottom = Inches(0)


def fill_cover(slide):
    sh = heading_shape(slide)
    if sh is None or not getattr(C, "COVER", None):
        return
    for p in sh.text_frame.paragraphs:
        txt = "".join(r.text for r in p.runs).strip()
        for label, value in C.COVER.items():
            if txt.lower().startswith(label.lower()) and value:
                r = p.add_run()
                r.text = " " + value
                src = p.runs[0].font
                r.font.name, r.font.size = C.FONT, src.size
                r.font.bold = True
                if src.color and src.color.type is not None:
                    r.font.color.rgb = src.color.rgb


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT / "out/deck.pptx"))
    a = ap.parse_args()
    prs = Presentation(TEMPLATE)
    slides = restructure(prs)
    for i, (n, slide) in enumerate(slides):
        key = getattr(C, "KEYS", [None] * len(slides))[i] if hasattr(C, "KEYS") else n
        if n == 1:
            fill_cover(slide)
        s = C.SLIDES.get(key)
        if not s:
            continue
        set_heading(slide, s.get("label"), move=s.get("move_label", True))
        if s.get("title"):
            add_text(slide, {"name": "ms-action-title", "x": C.CONTENT_X, "y": C.TITLE_Y, "w": C.CONTENT_W, "h": C.TITLE_H,
                             "paras": [s["title"]], "font": C.FONT_SEMI, "size": C.TITLE_PT, "color": C.INK_STRONG,
                             "anchor": "t", "line": 0.95})
        for e in s.get("elements", []):
            shape = ADD[e["type"]](slide, e)
            if not shape.name.startswith("ms-"):
                shape.name = "ms-" + shape.name   # marks builder shapes (check.py, heading_shape)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"deck   {out}  ({len(slides)} slides: {', '.join(str(n) for n, _ in slides)})")


if __name__ == "__main__":
    main()
